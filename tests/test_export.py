# -*- coding: utf-8 -*-
"""tests for _build_xlsx / _build_csv 导出辅助函数."""

from __future__ import annotations

import io
from datetime import datetime, timezone

import pytest

from javert.web.api.routes_workbench import _build_csv, _build_xlsx, _xl_cell


def test_xl_cell_none():
    assert _xl_cell(None) == ""


def test_xl_cell_datetime_naive_treated_utc():
    dt = datetime(2026, 5, 20, 14, 30, 22)
    out = _xl_cell(dt)
    assert out == "2026-05-20 14:30:22"


def test_xl_cell_datetime_with_tz():
    dt = datetime(2026, 5, 20, 14, 30, 22, tzinfo=timezone.utc)
    assert _xl_cell(dt) == "2026-05-20 14:30:22"


def test_xl_cell_bool():
    assert _xl_cell(True) == "是"
    assert _xl_cell(False) == "否"


def test_build_csv_utf8_bom():
    rows = [
        {"run_id": "aud_abc", "patient_id": "J66252", "reviewer": "alice"},
        {"run_id": "aud_def", "patient_id": "J18906", "reviewer": "bob"},
    ]
    data = _build_csv(rows)
    assert data.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    text = data[3:].decode("utf-8")
    assert "run_id,patient_id,reviewer" in text
    assert "alice" in text
    assert "bob" in text


def test_build_csv_empty():
    data = _build_csv([])
    assert data.startswith(b"\xef\xbb\xbf")
    assert data[3:] == b""


def test_build_csv_handles_datetime():
    rows = [
        {"id": 1, "ts": datetime(2026, 5, 20, 14, 30, 22, tzinfo=timezone.utc)},
    ]
    data = _build_csv(rows)
    text = data[3:].decode("utf-8")
    assert "2026-05-20 14:30:22" in text


def test_build_xlsx_multi_sheet():
    sheets = {
        "审核结果": [
            {"run_id": "aud_abc", "patient_id": "J66252", "reviewer": "alice"},
        ],
        "审核进度": [
            {"reviewer": "alice", "count": 5},
        ],
        "规则维度": [
            {"rule_id": "R191", "javert_v": 18},
        ],
    }
    data = _build_xlsx(sheets)
    # 是合法 zip-archive 起头
    assert data[:2] == b"PK"
    # 用 openpyxl 反序列化看 sheet 名 + 内容
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert set(wb.sheetnames) == {"审核结果", "审核进度", "规则维度"}
    ws = wb["审核结果"]
    headers = [c.value for c in ws[1]]
    assert headers == ["run_id", "patient_id", "reviewer"]
    row2 = [c.value for c in ws[2]]
    assert row2 == ["aud_abc", "J66252", "alice"]


def test_build_xlsx_empty_sheet():
    sheets = {"空表": []}
    data = _build_xlsx(sheets)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert "空表" in wb.sheetnames
    ws = wb["空表"]
    # 至少一行 (空) 占位
    assert ws[1][0].value == "(空)"


def test_build_xlsx_handles_long_sheet_name():
    # Excel sheet 名 ≤ 31 字符 限制 — _build_xlsx 截断
    long_name = "a" * 50
    sheets = {long_name: [{"x": 1}]}
    data = _build_xlsx(sheets)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert any(len(n) <= 31 for n in wb.sheetnames)
