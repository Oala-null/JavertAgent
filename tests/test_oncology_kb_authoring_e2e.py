from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from javert.data.loader import DataLoader
from javert.oncology.authoring.curated import build_manifest
from javert.oncology.authoring.ids import checksum
from javert.oncology.authoring.regimens import build_regimen_authoring_records
from javert.oncology.authoring.release import (
    AUTHORITY_SCHEMA_VERSION,
    DRUG_ASSET,
    ELIGIBILITY_ASSET,
    PATHOLOGY_ASSET,
    REGIMEN_ASSET,
    CoverageItem,
    CuratedAtomCoverage,
    ReleaseRevision,
    activate_release,
    build_release_authority_snapshot,
    build_release_candidate,
    compile_release,
    publish_candidate,
    resolve_active_release_assets,
    validate_compiled_release,
    write_release_bundle,
)
from javert.oncology.authoring.sources import (
    build_authoring_snapshot,
    build_candidate_universe,
)
from javert.oncology.authoring.staging import InMemoryKnowledgeStore
from javert.oncology.authoring.workbooks import parse_workbook, validate_workbook
from javert.tools.drug_audit_lookup import lookup_patient_drugs


ROOT = Path(__file__).resolve().parents[1]
AUTHORING_OUTPUT = ROOT / "outputs/add-oncology-kb-authoring"
ELIGIBILITY_WORKBOOK = AUTHORING_OUTPUT / "肿瘤药指南适应证与医保限定条件树KB.xlsx"
REGIMEN_WORKBOOK = AUTHORING_OUTPUT / "肿瘤治疗方案组成KB.xlsx"
SOURCE_CHECKSUM = "sha256:" + "1" * 64
SOURCE_REF = {
    "source_id": "source:offline-e2e",
    "title": "离线端到端合成来源",
    "version": "1.0",
    "publication_date": "2026-01-01",
    "effective_date": None,
    "retrieval_date": "2026-07-21",
    "checksum": SOURCE_CHECKSUM,
}
PATIENT_ID = "E2E-ONCOLOGY-AUTHORING"
DRUG_NAME = "合成肿瘤药"
DRUG_CONCEPT_ID = "drug-synthetic-oncology"
DRUG_CODE = "SYN-ONC-001"


class _Loader(DataLoader):
    def __init__(self, fees: pd.DataFrame):
        self._fees = fees

    def get_fees(self, patient_id: str) -> pd.DataFrame:
        return self._fees

    def get_notes(self, patient_id: str) -> pd.DataFrame:
        return pd.DataFrame()

    def all_fees(self) -> pd.DataFrame:
        return self._fees

    def all_notes(self) -> pd.DataFrame:
        return pd.DataFrame()


def _copy_and_approve_first_review(
    source: Path,
    target: Path,
    *,
    review_sheet: str,
) -> Path:
    shutil.copy2(source, target)
    workbook = load_workbook(target)
    sheet = workbook[review_sheet]
    columns = {str(cell.value): cell.column for cell in sheet[1]}
    values = {
        "review_decision": "APPROVE",
        "expert_comment": "离线端到端审核通过",
        "reviewer_id": "domain-reviewer-e2e",
        "reviewed_at": "2026-07-21T12:00:00Z",
    }
    for name, value in values.items():
        sheet.cell(row=2, column=columns[name], value=value)
    workbook.save(target)
    return target


def _release_revision(
    *,
    asset_name: str,
    revision_id: str,
    logical_id: str,
    entity_id: str,
    payload: dict,
    policy_scope: str = "",
) -> ReleaseRevision:
    return ReleaseRevision(
        revision_id=revision_id,
        logical_id=logical_id,
        asset_name=asset_name,
        entity_id=entity_id,
        payload=payload,
        reviewer_id="domain-reviewer-e2e",
        effective_from=date(2026, 1, 1),
        effective_to=date(2027, 12, 31),
        source_refs=(SOURCE_REF,),
        policy_scope=policy_scope,
        source_item_ids=(f"source-{revision_id}",),
    )


def _approved_revisions() -> list[ReleaseRevision]:
    diagnosis_leaf = {
        "node_id": "node-synthetic-diagnosis",
        "kind": "leaf",
        "source_text": "合成癌诊断",
        "children": [],
        "criterion_id": "criterion-synthetic-diagnosis",
        "criterion_type": "diagnosis",
        "expected": {"includes": ["合成癌"]},
        "evidence_policy": {"anchored": True, "missing_is": "UNKNOWN"},
        "documentation_template": "",
    }
    drug_entries = [
        {
            "rule_type": "限适应症",
            "source_type": "insurance",
            "basis": "医保支付限定：合成癌",
            "source_refs": ["source:offline-e2e:insurance"],
        },
        {
            "rule_type": "限适应症",
            "source_type": "guideline",
            "basis": "指南适应证：合成癌",
            "source_refs": ["source:offline-e2e:guideline"],
        },
        {
            "rule_type": "超说明书",
            "source_type": "label",
            "basis": "合成说明书安全边界",
            "source_refs": ["source:offline-e2e:label"],
        },
        {
            "rule_type": "限二线",
            "source_type": "line_of_therapy",
            "basis": "合成二线限定",
            "source_refs": ["source:offline-e2e:line"],
        },
        {
            "rule_type": "禁忌症",
            "source_type": "safety",
            "basis": "合成独立安全禁忌",
            "source_refs": ["source:offline-e2e:safety"],
        },
    ]
    return [
        _release_revision(
            asset_name=DRUG_ASSET,
            revision_id="rev-drug-e2e",
            logical_id="drug-logical-e2e",
            entity_id=DRUG_NAME,
            payload={
                "canonical_name": DRUG_NAME,
                "canonical_match_name": DRUG_NAME,
                "aliases": ["Synthetic Oncology Drug"],
                "codes": [DRUG_CODE],
                "oncology": {
                    "drug_concept_id": DRUG_CONCEPT_ID,
                    "name_fallback": "exact_entity_name",
                },
                "entries": drug_entries,
                "insurance_status": "restricted",
                "sources": {},
            },
        ),
        _release_revision(
            asset_name=ELIGIBILITY_ASSET,
            revision_id="rev-eligibility-insurance-e2e",
            logical_id="eligibility-insurance-e2e",
            entity_id="eligibility-insurance-e2e",
            policy_scope="INSURANCE_PAYMENT",
            payload={
                "rule_id": "eligibility-insurance-e2e",
                "drug_concept_id": DRUG_CONCEPT_ID,
                "indication_branch_id": "branch-insurance-e2e",
                "version": "1.0.0",
                "raw_restriction": "医保支付限定：合成癌",
                "condition_tree": diagnosis_leaf,
            },
        ),
        _release_revision(
            asset_name=ELIGIBILITY_ASSET,
            revision_id="rev-eligibility-guideline-e2e",
            logical_id="eligibility-guideline-e2e",
            entity_id="eligibility-guideline-e2e",
            policy_scope="GUIDELINE_INDICATION",
            payload={
                "rule_id": "eligibility-guideline-e2e",
                "drug_concept_id": DRUG_CONCEPT_ID,
                "indication_branch_id": "branch-guideline-e2e",
                "version": "1.0.0",
                "raw_restriction": "指南适应证：合成癌",
                "condition_tree": diagnosis_leaf,
            },
        ),
        _release_revision(
            asset_name=PATHOLOGY_ASSET,
            revision_id="rev-pathology-e2e",
            logical_id="pathology-e2e",
            entity_id="pathology-e2e",
            payload={
                "entry_id": "pathology-e2e",
                "marker_id": "SYN1",
                "aliases": [{"value": "SYN1", "alias_type": "gene"}],
                "observation_type": "sequence_variant",
                "methods": ["MOLECULAR"],
                "cancer_contexts": ["合成癌"],
                "policy_context": "offline-e2e",
                "specimen_constraints": [],
                "scoring_system": "",
                "accepted_values": ["positive"],
                "threshold": {"operator": "EQUALS", "value": "positive"},
                "raw_condition": "SYN1 阳性",
            },
        ),
        _release_revision(
            asset_name=REGIMEN_ASSET,
            revision_id="rev-regimen-e2e",
            logical_id="regimen-e2e",
            entity_id="regimen-e2e",
            payload={
                "regimen_id": "regimen-e2e",
                "canonical_name": "SYN-E2E",
                "aliases": ["SYN-E2E"],
                "cancer_contexts": ["合成癌"],
                "components": [
                    {"drug_concept_id": DRUG_CONCEPT_ID, "token": "SYN"}
                ],
            },
        ),
    ]


def _coverage(revisions: list[ReleaseRevision]) -> list[CoverageItem]:
    return [
        CoverageItem(
            source_item_id=revision.source_item_ids[0],
            source_type="SYNTHETIC",
            disposition="approved",
            revision_id=revision.revision_id,
            source_checksum=SOURCE_CHECKSUM,
        )
        for revision in revisions
    ]


def _lookup(
    *,
    loader: _Loader,
    drug_path: Path,
    eligibility_path: Path,
    pathology_path: Path,
    regimen_path: Path,
    zd_path: Path,
    rule_id: str,
    rule_type: str,
) -> dict:
    return lookup_patient_drugs(
        PATIENT_ID,
        loader,
        drug_path,
        zd_path,
        rule_type=rule_type,
        source_type="insurance" if rule_id == "RD04" else None,
        oncology_v2_mode="on",
        audit_rule_id=rule_id,
        eligibility_path=eligibility_path,
        pathology_path=pathology_path,
        regimen_path=regimen_path,
    )


def test_complete_offline_authoring_to_rd04_dual_scope_evaluation(tmp_path: Path) -> None:
    # 五类来源并集和精选知识原子均从当前仓库源文件实时构建。
    universe = build_candidate_universe(ROOT)
    five_source_categories = {
        "hospital_catalog": lambda row: row.source_membership["hospital_catalog"],
        "national_catalog": lambda row: row.source_membership["national_catalog"],
        "guideline": lambda row: row.source_membership["guideline"],
        "existing_json_assets": lambda row: (
            row.source_membership["oncology_drug_asset"]
            or row.source_membership["structured_asset"]
        ),
        "rule_yaml": lambda row: row.source_membership["rule_yaml"],
    }
    assert universe
    assert all(any(predicate(row) for row in universe) for predicate in five_source_categories.values())
    curated = build_manifest(ROOT)
    assert curated["counts"]["rules"] == 28
    assert curated["atoms"]

    # 两份现成 Excel 必须与当前候选、精选原子和方案投影保持可追溯关系。
    assert validate_workbook(ELIGIBILITY_WORKBOOK, kind="eligibility") == []
    assert validate_workbook(REGIMEN_WORKBOOK, kind="regimen") == []
    eligibility_export = parse_workbook(ELIGIBILITY_WORKBOOK, kind="eligibility")
    regimen_export = parse_workbook(REGIMEN_WORKBOOK, kind="regimen")
    snapshot = build_authoring_snapshot(ROOT)
    exported_fragments = {
        row["source_fragment_id"]
        for row in eligibility_export["sheets"]["03_来源原文"]
    }
    exported_atoms = {
        row["atom_id"] for row in eligibility_export["sheets"]["09_肿瘤知识保全"]
    }
    assert {item.source_fragment_id for item in snapshot.source_fragments} <= exported_fragments
    assert {item["atom_id"] for item in curated["atoms"]} <= exported_atoms
    regimen_records = build_regimen_authoring_records(ROOT)
    exported_regimen_revisions = {
        row["regimen_revision_id"]
        for row in regimen_export["sheets"]["02_方案主表"]
    }
    assert {
        item.regimen_revision_id for item in regimen_records["revisions"]
    } <= exported_regimen_revisions

    # 在 tmp 副本中模拟具名专家的合法审核，不改机器列和 checksum。
    reviewed_eligibility = _copy_and_approve_first_review(
        ELIGIBILITY_WORKBOOK,
        tmp_path / ELIGIBILITY_WORKBOOK.name,
        review_sheet="06_专家审核",
    )
    reviewed_regimen = _copy_and_approve_first_review(
        REGIMEN_WORKBOOK,
        tmp_path / REGIMEN_WORKBOOK.name,
        review_sheet="07_专家审核",
    )
    assert validate_workbook(reviewed_eligibility, kind="eligibility") == []
    assert validate_workbook(reviewed_regimen, kind="regimen") == []

    # 回导后依次走 upload、服务端二次校验和事务性 materialize。
    store = InMemoryKnowledgeStore()
    materialized_batches = []
    for workbook, kind in (
        (reviewed_eligibility, "eligibility"),
        (reviewed_regimen, "regimen"),
    ):
        batch = store.upload(workbook, kind=kind, uploaded_by="authoring-operator-e2e")
        assert batch.status == "UPLOADED"
        store.server_validate(batch.import_batch_id)
        assert batch.status == "VALIDATED"
        store.materialize(batch.import_batch_id)
        assert batch.status == "MATERIALIZED"
        assert sum(item["staged"] for item in batch.reconciliation.values()) == len(batch.rows)
        materialized_batches.append(batch)
    assert any(
        row.get("review_decision") == "APPROVE"
        for rows in store.authoring.values()
        for row in rows.values()
    )

    # 用现有 release 合同构造经批准、职责分离的四资产双 scope 发布。
    revisions = _approved_revisions()
    coverage = _coverage(revisions)
    release_atoms = [
        CuratedAtomCoverage(
            atom_id=atom["atom_id"],
            source_rule_id=atom["source_rule_id"],
            migration_status="VERIFIED" if atom["oncology"] else atom["migration_status"],
            oncology=atom["oncology"],
            rule_status=atom["rule_status"],
            target_id=atom["target_id"],
            verification_evidence=(
                "offline-e2e:curated-oncology" if atom["oncology"] else ""
            ),
        )
        for atom in curated["atoms"]
    ]
    source_payload = {
        "schema_version": AUTHORITY_SCHEMA_VERSION,
        "source_item_ids": sorted(item.source_item_id for item in coverage),
        "counts": {"source_items": len(coverage)},
        "candidate_universe_checksum": checksum(
            [row.candidate_id for row in universe]
        ),
    }
    source_manifest = {
        **source_payload,
        "snapshot_checksum": checksum(source_payload),
    }
    curated_payload = {
        "schema_version": AUTHORITY_SCHEMA_VERSION,
        "atoms": [
            {
                "atom_id": atom.atom_id,
                "source_rule_id": atom.source_rule_id,
                "migration_status": atom.migration_status,
                "oncology": atom.oncology,
                "rule_status": atom.rule_status,
                "target_id": atom.target_id,
                "verification_evidence": atom.verification_evidence,
            }
            for atom in release_atoms
        ],
        "counts": {"atoms": len(release_atoms)},
    }
    curated_authority_manifest = {
        **curated_payload,
        "manifest_checksum": checksum(curated_payload),
    }
    authority = build_release_authority_snapshot(
        source_snapshot_manifest=source_manifest,
        curated_manifest=curated_authority_manifest,
        expected_source_snapshot_checksum=source_manifest["snapshot_checksum"],
        expected_curated_manifest_checksum=curated_authority_manifest["manifest_checksum"],
    )
    candidate = build_release_candidate(
        revisions=revisions,
        coverage=coverage,
        curated_atoms=release_atoms,
        source_snapshot_checksum=authority.source_snapshot_checksum,
        created_by="release-builder-e2e",
        release_operator="release-operator-e2e",
        created_at="2026-07-21T13:00:00Z",
        expected_source_item_ids=authority.source_item_ids,
        authority_snapshot=authority,
        asset_extras={
            REGIMEN_ASSET: {
                "drug_concepts": [
                    {
                        "concept_id": DRUG_CONCEPT_ID,
                        "generic_name": DRUG_NAME,
                        "aliases": [
                            {
                                "value": "Synthetic Oncology Drug",
                                "alias_type": "english_generic",
                            }
                        ],
                        "insurance_codes": [DRUG_CODE],
                    }
                ]
            },
            DRUG_ASSET: {
                "policy": {"guideline_authority": "临床指导原则，不冒充法定说明书"},
                "sources": {"offline_e2e": SOURCE_REF},
            },
        },
    )
    candidate_compiled = compile_release(candidate)
    validate_compiled_release(candidate_compiled, require_published=False)
    assert candidate_compiled.manifest["release_status"] == "CANDIDATE"
    published = publish_candidate(
        candidate,
        published_by="release-operator-e2e",
        published_at="2026-07-21T14:00:00Z",
    )
    compiled = compile_release(published)
    validate_compiled_release(compiled)
    for revision in revisions:
        store.approve(revision.revision_id)
    store.register_release(compiled.release_id)
    release_dir = tmp_path / "releases"
    write_release_bundle(compiled, release_dir)
    activate_release(release_dir, compiled.release_id)
    assets = resolve_active_release_assets(release_dir)
    assert set(assets) == {DRUG_ASSET, ELIGIBILITY_ASSET, PATHOLOGY_ASSET, REGIMEN_ASSET}
    store.deploy(compiled.release_id)
    assert store.deployed_release_id == compiled.release_id

    # active release 真实进入 RD04 求值；同一药品只生成两个独立 scope 状态。
    fees = pd.DataFrame(
        {
            "bah": [f"H-{PATIENT_ID}"],
            "fee_ocur_time": ["2026-06-15"],
            "cnt": [1],
            "medins_list_name": [DRUG_NAME],
            "medins_chrgitm_type": ["西药"],
            "med_list_codg": [DRUG_CODE],
        }
    )
    loader = _Loader(fees)
    zd_path = tmp_path / "synthetic-zd.csv"
    pd.DataFrame(
        {
            "ba_id": [f"H-{PATIENT_ID}"],
            "diag_name": ["合成癌"],
            "diag_code": ["SYN-C80"],
            "maindiag_flag": ["1"],
        }
    ).to_csv(zd_path, index=False)
    lookup_args = {
        "loader": loader,
        "drug_path": assets[DRUG_ASSET],
        "eligibility_path": assets[ELIGIBILITY_ASSET],
        "pathology_path": assets[PATHOLOGY_ASSET],
        "regimen_path": assets[REGIMEN_ASSET],
        "zd_path": zd_path,
    }
    rd04 = _lookup(**lookup_args, rule_id="RD04", rule_type="限适应症")
    assert "error" not in rd04["oncology_structured"]
    assert len(rd04["matches"]) == 2
    candidate_evaluation = rd04["oncology_structured"]["candidate_evaluations"][0]
    evaluations = candidate_evaluation["eligibility_evaluations"]
    assert {item["policy_scope"] for item in evaluations} == {
        "INSURANCE_PAYMENT",
        "GUIDELINE_INDICATION",
    }
    assert len(evaluations) == 2
    assert all(item["release_id"] == compiled.release_id for item in evaluations)
    assert all(item["eligibility_status"] == "SATISFIED" for item in evaluations)

    # 肿瘤资格只归 RD04；通用 bulk 排除，RD03 仍保留独立安全语义且 key 不冲突。
    rd04_keys = {item["ownership_key"] for item in rd04["matches"]}
    for rule_id, rule_type in (
        ("R007", "限适应症"),
        ("RD01", "超说明书"),
        ("RD02", "限二线"),
    ):
        general = _lookup(**lookup_args, rule_id=rule_id, rule_type=rule_type)
        assert general["matches"] == []
    rd03 = _lookup(**lookup_args, rule_id="RD03", rule_type="禁忌症")
    assert {item["rule_type"] for item in rd03["matches"]} == {"禁忌症"}
    assert not (rd04_keys & {item["ownership_key"] for item in rd03["matches"]})
