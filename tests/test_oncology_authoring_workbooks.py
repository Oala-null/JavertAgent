from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from javert.oncology.authoring.ids import checksum
from javert.oncology.authoring.workbooks import (
    ELIGIBILITY_SHEETS,
    EXPERT_COLUMNS,
    REGIMEN_SHEETS,
    canonical_workbook_payload,
    parse_workbook,
    validate_workbook,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs/add-oncology-kb-authoring"
ELIGIBILITY = OUTPUT / "肿瘤药指南适应证与医保限定条件树KB.xlsx"
REGIMEN = OUTPUT / "肿瘤治疗方案组成KB.xlsx"


def test_eligibility_workbook_schema_validation_and_protection() -> None:
    workbook = load_workbook(ELIGIBILITY)
    assert tuple(workbook.sheetnames) == ELIGIBILITY_SHEETS
    assert not validate_workbook(ELIGIBILITY, kind="eligibility")
    sheet = workbook["04_适应证分支"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.freeze_panes == "C2"
    assert sheet.protection.sheet
    assert sheet.cell(2, headers["branch_id"]).protection.locked
    assert not sheet.cell(2, headers["review_decision"]).protection.locked
    assert sheet.data_validations.count > 0
    preservation = workbook["09_肿瘤知识保全"]
    preservation_headers = {
        cell.value: cell.column for cell in preservation[1]
    }
    assert preservation.cell(
        2, preservation_headers["target_id"]
    ).protection.locked
    assert not preservation.cell(
        2, preservation_headers["expert_target_id"]
    ).protection.locked
    review = workbook["06_专家审核"]
    review_headers = {cell.value: cell.column for cell in review[1]}
    assert any(
        review.cell(row, review_headers["entity_type"]).value == "drug_class"
        for row in range(2, review.max_row + 1)
    )


def test_regimen_workbook_schema_validation_and_reserved_schedule() -> None:
    workbook = load_workbook(REGIMEN)
    assert tuple(workbook.sheetnames) == REGIMEN_SHEETS
    assert not validate_workbook(REGIMEN, kind="regimen")
    schedule = workbook["06_预留给药字段"]
    headers = {cell.value: cell.column for cell in schedule[1]}
    assert all(schedule.cell(row, headers["publishing_enabled"]).value is False for row in range(2, schedule.max_row + 1))
    assert all(schedule.cell(row, headers["inference_enabled"]).value is False for row in range(2, schedule.max_row + 1))
    assert schedule.protection.sheet
    assert not schedule.cell(2, headers["dose_value"]).protection.locked


def test_generated_workbooks_have_stable_ids_checksums_and_dates() -> None:
    eligibility = parse_workbook(ELIGIBILITY, kind="eligibility")
    branches = eligibility["sheets"]["04_适应证分支"]
    assert all(row["branch_id"] and row["row_checksum"] for row in branches)
    assert all(row["effective_from"] == "2026-01-01" for row in branches)
    assert all(row["effective_to"] == "2027-12-31" for row in branches)
    assert all(row["effective_date_basis"] == "CURRENT_FILE_ASSUMPTION" for row in branches)


def test_curated_authority_fields_round_trip_and_are_checksum_protected(
    tmp_path: Path,
) -> None:
    manifest = json.loads(
        (ROOT / "docs/oncology/authoring/curated_knowledge_manifest.json").read_text(
            encoding="utf-8"
        )
    )
    expected = {atom["atom_id"]: atom for atom in manifest["atoms"]}
    parsed = parse_workbook(ELIGIBILITY, kind="eligibility")
    candidates = json.loads(
        (ROOT / "docs/oncology/authoring/oncology_authoring_candidates.json").read_text(
            encoding="utf-8"
        )
    )
    expected_snapshot = checksum(
        {
            "candidate_snapshot_checksum": candidates["snapshot_checksum"],
            "curated_manifest_checksum": manifest["manifest_checksum"],
        }
    )
    assert parsed["metadata"]["candidate_snapshot_checksum"] == candidates["snapshot_checksum"]
    assert parsed["metadata"]["curated_manifest_checksum"] == manifest["manifest_checksum"]
    assert parsed["metadata"]["source_snapshot_checksum"] == expected_snapshot
    rows = parsed["sheets"]["09_肿瘤知识保全"]
    assert len(rows) == len(expected)
    for row in rows:
        atom = expected[row["atom_id"]]
        assert type(row["oncology"]) is bool
        assert row["oncology"] is atom["oncology"]
        assert row["rule_status"] == atom["rule_status"]

    changed = tmp_path / ELIGIBILITY.name
    workbook = load_workbook(ELIGIBILITY)
    sheet = workbook["09_肿瘤知识保全"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["oncology"]).value = not sheet.cell(
        2, headers["oncology"]
    ).value
    workbook.save(changed)
    assert any(
        issue.error_code == "ROW_CHECKSUM_MISMATCH"
        and issue.sheet == "09_肿瘤知识保全"
        for issue in validate_workbook(changed, kind="eligibility")
    )


def test_canonical_round_trip_is_order_and_style_independent(tmp_path: Path) -> None:
    for source, kind in ((ELIGIBILITY, "eligibility"), (REGIMEN, "regimen")):
        before = canonical_workbook_payload(source, kind=kind)
        workbook = load_workbook(source)
        copy = tmp_path / source.name
        workbook.save(copy)
        after = canonical_workbook_payload(copy, kind=kind)
        assert before == after


def test_workbooks_contain_no_sensitive_patterns() -> None:
    for source, kind in ((ELIGIBILITY, "eligibility"), (REGIMEN, "regimen")):
        parsed = parse_workbook(source, kind=kind)
        dumped = str(parsed)
        for forbidden in ("patient_id", "ownership_id", "run_id", "/Users/", "PWD="):
            assert forbidden not in dumped
