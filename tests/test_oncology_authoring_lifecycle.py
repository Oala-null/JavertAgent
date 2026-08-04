from __future__ import annotations

from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from javert.oncology.authoring.lifecycle import (
    ReviewLedger,
    RevisionRecord,
    clone_superseding_revision,
    edit_revision,
    review_event_id,
    transition_revision,
    validate_effective_window_overlaps,
)
from javert.oncology.authoring.models import ReviewDecision, ReviewEvent, RevisionLifecycle
from javert.oncology.authoring.workbooks import validate_workbook


ROOT = Path(__file__).resolve().parents[1]
ELIGIBILITY = ROOT / "outputs/add-oncology-kb-authoring/肿瘤药指南适应证与医保限定条件树KB.xlsx"


def _revision(lifecycle: RevisionLifecycle = RevisionLifecycle.DRAFT) -> RevisionRecord:
    return RevisionRecord(
        logical_id="rule_synthetic",
        revision_id="rev_synthetic_1",
        lifecycle=lifecycle,
        content={"value": 1},
        effective_from=date(2026, 1, 1),
        effective_to=date(2027, 12, 31),
    )


def test_revision_state_machine_and_immutable_clone() -> None:
    review = transition_revision(_revision(), RevisionLifecycle.IN_REVIEW)
    approved = transition_revision(review, RevisionLifecycle.APPROVED)
    with pytest.raises(ValueError, match="不可原地修改"):
        edit_revision(approved, {"value": 2})
    cloned = clone_superseding_revision(approved, {"value": 2})
    assert cloned.lifecycle == RevisionLifecycle.DRAFT
    assert cloned.supersedes_revision_id == approved.revision_id
    assert cloned.revision_id != approved.revision_id


def test_inclusive_overlap_detects_versions_but_not_shared_or_branches() -> None:
    rows = [
        {"logical_id": "r", "revision_id": "v1", "lifecycle": "APPROVED", "effective_from": "2026-01-01", "effective_to": "2026-12-31", "branch_id": "b1"},
        {"logical_id": "r", "revision_id": "v1", "lifecycle": "APPROVED", "effective_from": "2026-01-01", "effective_to": "2026-12-31", "branch_id": "b2"},
        {"logical_id": "r", "revision_id": "v2", "lifecycle": "APPROVED", "effective_from": "2026-12-31", "effective_to": "2027-12-31", "branch_id": "b3"},
    ]
    assert validate_effective_window_overlaps(rows) == [("v1", "v2")]
    rows[2]["effective_from"] = "2027-01-01"
    assert not validate_effective_window_overlaps(rows)


def test_review_events_are_append_only_and_projection_keeps_history() -> None:
    ledger = ReviewLedger()
    at1 = "2026-07-21T01:00:00+00:00"
    first_id = review_event_id("entity", ReviewDecision.UNABLE_TO_DETERMINE, at1, "reviewer-a")
    first = ReviewEvent(
        event_id=first_id,
        entity_type="condition_node",
        entity_id="entity",
        field_name="operator",
        decision=ReviewDecision.UNABLE_TO_DETERMINE,
        reviewer_id="reviewer-a",
        reviewed_at=datetime.fromisoformat(at1),
        comment="证据不足",
    )
    ledger.append(first)
    second = ReviewEvent(
        event_id=review_event_id("entity", ReviewDecision.APPROVE, "2026-07-21T02:00:00+00:00", "reviewer-b"),
        entity_type="condition_node",
        entity_id="entity",
        field_name="operator",
        decision=ReviewDecision.APPROVE,
        reviewer_id="reviewer-b",
        reviewed_at=datetime(2026, 7, 21, 2, tzinfo=timezone.utc),
        comment="已补证",
        previous_event_id=first_id,
    )
    ledger.append(second)
    projection = ledger.project("entity")
    assert projection["decision"] == "APPROVE"
    assert projection["event_count"] == 2
    assert [item["decision"] for item in projection["history"]] == ["UNABLE_TO_DETERMINE", "APPROVE"]


def test_missing_required_column_fails_closed_with_safe_diagnostic(tmp_path: Path) -> None:
    workbook = load_workbook(ELIGIBILITY)
    sheet = workbook["05_条件节点"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.delete_cols(headers["operator"])
    broken = tmp_path / "broken.xlsx"
    workbook.save(broken)
    issues = validate_workbook(broken, kind="eligibility")
    assert [(issue.error_code, issue.sheet) for issue in issues] == [("COLUMN_MISSING", "05_条件节点")]
    assert all("original_text" not in issue.safe_detail for issue in issues)


def test_date_override_on_approved_revision_is_rejected(tmp_path: Path) -> None:
    workbook = load_workbook(ELIGIBILITY)
    sheet = workbook["04_适应证分支"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["effective_from"], "2026-02-01")
    sheet.cell(2, headers["effective_date_basis"], "EXPERT_OVERRIDE")
    sheet.cell(2, headers["date_override_reason"], "合成覆盖原因")
    sheet.cell(2, headers["date_review_comment"], "合成日期审核")
    sheet.cell(2, headers["lifecycle"], "APPROVED")
    broken = tmp_path / "immutable-date.xlsx"
    workbook.save(broken)
    codes = {issue.error_code for issue in validate_workbook(broken, kind="eligibility")}
    assert "IMMUTABLE_DATE_EDIT" in codes
    assert "ROW_CHECKSUM_MISMATCH" in codes
