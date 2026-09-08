#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从已研究反馈表构建/校验慢病 draft 条件资产。"""

from __future__ import annotations

import argparse
import json
import re
import sys
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from javert.chronic.contracts import (  # noqa: E402
    ChronicDiseaseCriteriaAsset,
    ExpertInterpretation,
    SourceManifest,
)
from javert.chronic.knowledge import (  # noqa: E402
    asset_payload_checksum,
    canonical_json_bytes,
    fragment_payload_checksum,
    interpretation_payload_checksum,
    load_criteria_asset,
    load_source_manifest,
    source_manifest_payload_checksum,
    validate_source_isolation,
    sha256_digest,
)


SOURCES_OUT = ROOT / "configs" / "chronic_disease_sources.json"
CRITERIA_OUT = ROOT / "configs" / "chronic_disease_criteria.json"
SCHEMA_OUT = ROOT / "configs" / "schemas" / "chronic_disease_criteria.schema.json"
SOURCE_SCHEMA_OUT = ROOT / "configs" / "schemas" / "chronic_disease_sources.schema.json"
FEEDBACK_OUT = ROOT / "configs" / "chronic_disease_expert_feedback.json"
DRAFT_SOURCES = ROOT / "configs" / "chronic_disease_sources.draft-r1.json"
DRAFT_CRITERIA = ROOT / "configs" / "chronic_disease_criteria.draft-r1.json"
# r2只编译这次用户明确转录的五条解释；新解释须创建下一revision。
R2_INTERPRETATIONS = [
    ("CD07", "H4", "高血压3级or（高血压1-2级+靶器官损害/临床综合征）"),
    ("CD19", "H5", "总分需≥6分andX线为III期+"),
    ("CD09", "H6", "病史是必备条件，病史and（HBV/HCV分支）"),
    ("CD01", "H7", "CBCand骨髓穿刺and骨髓活检and排除检查"),
    ("CD03", "H8", "糖尿病肾病4个子项满足任一即可"),
]

CURRENT_SOURCE_ID = "hlj-outpatient-chronic-2025"
HISTORICAL_SOURCE_ID = "heihe-outpatient-chronic-2020"
CURRENT_TITLE = "黑龙江省基本医疗保险门诊慢性病认定标准（试行）"
CURRENT_NUMBER = "黑市医保发〔2025〕50号 / 黑医保规〔2025〕12号"
HISTORICAL_TITLE = "黑河市基本医疗保险门诊慢性病鉴定标准"
HISTORICAL_NUMBER = "黑市医保发〔2020〕37号"

DISEASE_IDS = {
    "CD01": "aplastic-anemia",
    "CD02": "chronic-kidney-disease-stage-3-plus",
    "CD03": "diabetes-with-complications",
    "CD04": "atrial-fibrillation",
    "CD05": "coronary-heart-disease-nyha-3-plus",
    "CD06": "rheumatic-heart-disease-nyha-3-plus",
    "CD07": "hypertension-grade-3-plus",
    "CD08": "decompensated-cirrhosis",
    "CD09": "chronic-viral-hepatitis",
    "CD10": "hiv-aids",
    "CD11": "brucellosis",
    "CD12": "parkinson-disease",
    "CD13": "myasthenia-gravis",
    "CD14": "epilepsy",
    "CD15": "cerebrovascular-sequelae-with-limb-dysfunction",
    "CD16": "chronic-obstructive-pulmonary-disease",
    "CD17": "chronic-cor-pulmonale-heart-failure",
    "CD18": "bronchial-asthma",
    "CD19": "rheumatoid-arthritis-severe-limb-dysfunction",
    "CD20": "systemic-lupus-erythematosus",
}
BLOCKERS = {
    "CD01": {
        "block_reason_code": "DIAGNOSTIC_BASIS_COMBINATION_AMBIGUITY",
        "unresolved_question": "诊断依据1-4是全部必备、至少若干项还是其他组合？",
        "affected": ["ROOT", "CBC_2OF3", "MARROW_ASP", "MARROW_BIOPSY", "EXCLUSION"],
        "roles": ["医保办", "血液科专家"],
    },
    "CD07": {
        "block_reason_code": "POLICY_ROOT_LOGIC_CONFLICT",
        "unresolved_question": "3级血压路径与1-2级伴靶器官损害路径的根关系是OR还是AND？",
        "affected": ["ROOT", "BP_GRADE3", "GRADE12_DAMAGE"],
        "roles": ["医保办", "心血管专家"],
    },
    "CD09": {
        "block_reason_code": "POLICY_ROOT_CLOSURE_MISSING",
        "unresolved_question": "慢性病史是否必备，HBV/HCV分支如何与病史组合？",
        "affected": ["ROOT", "HBV_ANY", "HCV_ANY"],
        "roles": ["医保办", "感染或肝病专家"],
    },
    "CD19": {
        "block_reason_code": "SCORE_AND_FUNCTION_THRESHOLD_MISSING",
        "unresolved_question": "ACR/EULAR评分阈值、严重功能障碍分级及其与X线III期的组合是什么？",
        "affected": ["ROOT", "SCORE", "XRAY"],
        "roles": ["医保办", "风湿免疫专家"],
    },
}


def pretty_json_bytes(value: Any) -> bytes:
    return (
        json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n"
    ).encode("utf-8")


def _feedback_rows(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    overview = workbook["规则可行性"]
    node_sheet = workbook["规则节点明细"]
    disease_headers = [cell.value for cell in overview[7]]
    node_headers = [cell.value for cell in node_sheet[3]]
    diseases = [
        dict(zip(disease_headers, row, strict=True))
        for row in overview.iter_rows(min_row=8, max_row=27, values_only=True)
    ]
    nodes = [
        dict(zip(node_headers, row, strict=True))
        for row in node_sheet.iter_rows(min_row=4, max_row=61, values_only=True)
    ]
    if [row["规则ID（暂定）"] for row in diseases] != list(DISEASE_IDS):
        raise ValueError("反馈 Excel 必须按序恰好包含 CD01-CD20")
    if len(nodes) != 58:
        raise ValueError(f"反馈 Excel 主要节点必须为 58，实际 {len(nodes)}")
    return diseases, nodes


def _page_numbers(value: str) -> list[int]:
    match = re.fullmatch(r"P(\d+)(?:-P?(\d+))?", value.strip())
    if match is None:
        raise ValueError(f"无法解析 PDF 物理页: {value!r}")
    start, end = int(match.group(1)), int(match.group(2) or match.group(1))
    return list(range(start, end + 1))


def _fragment_id(rule_id: str, node_id: str, page: int) -> str:
    clean_node = re.sub(r"[^A-Za-z0-9_]+", "_", node_id).strip("_")
    return f"SRC2025-{rule_id}-{clean_node}-P{page:02d}"


def _document_rows() -> list[dict[str, Any]]:
    return [
        {
            "source_document_id": CURRENT_SOURCE_ID,
            "policy_version": "2025",
            "policy_role": "current_recognition",
            "title": CURRENT_TITLE,
            "document_number": CURRENT_NUMBER,
            "publication_date": "2025-09-02",
            "effective_from": "2025-09-02",
            "effective_to": None,
            "validity_text": "本通知自印发之日起实施，有效期2年；结束日待法务口径确认",
            "file_path": "慢病鉴定标准/黑河市医保局转发黑龙江省门诊慢性病认定标准通知(1).pdf",
            "physical_page_count": 19,
            "pdf_checksum": "sha256:" + "0" * 64,
        },
        {
            "source_document_id": HISTORICAL_SOURCE_ID,
            "policy_version": "2020",
            "policy_role": "historical_reference",
            "title": HISTORICAL_TITLE,
            "document_number": HISTORICAL_NUMBER,
            "publication_date": "2020-05-15",
            "effective_from": "2020-05-15",
            "effective_to": None,
            "validity_text": "历史鉴定标准，仅用于版本差异和授权历史重建",
            "file_path": "慢病鉴定标准/黑市医保发【2020】37号，关于印发《黑河市基本医疗保险门诊慢性病鉴定标准》的通知.pdf",
            "physical_page_count": 32,
            "pdf_checksum": "sha256:" + "0" * 64,
        },
    ]


def build_source_manifest(node_rows: list[dict[str, Any]]) -> dict[str, Any]:
    fragments: list[dict[str, Any]] = [
        {
            "source_fragment_id": "SRC2020-VERSION-BOUNDARY-P01",
            "source_document_id": HISTORICAL_SOURCE_ID,
            "document_title": HISTORICAL_TITLE,
            "document_number": HISTORICAL_NUMBER,
            "policy_version": "2020",
            "physical_page": 1,
            "printed_page_label": "1",
            "reviewed_excerpt": "关于印发《黑河市基本医疗保险门诊慢性病鉴定标准》的通知",
            "excerpt_kind": "verbatim",
            "extraction_method": "manual_visual_review",
            "review_status": "needs_review",
            "fragment_checksum": "sha256:" + "0" * 64,
        }
    ]
    for row in node_rows:
        threshold = str(row["运算符/阈值"] or "").strip()
        excerpt = str(row["节点条件"]).strip()
        if threshold:
            excerpt = f"{excerpt}；运算符/阈值：{threshold}"
        for page in _page_numbers(str(row["2025 PDF页"])):
            fragments.append(
                {
                    "source_fragment_id": _fragment_id(
                        str(row["规则ID"]), str(row["节点ID"]), page
                    ),
                    "source_document_id": CURRENT_SOURCE_ID,
                    "document_title": CURRENT_TITLE,
                    "document_number": CURRENT_NUMBER,
                    "policy_version": "2025",
                    "physical_page": page,
                    "printed_page_label": str(page - 1),
                    "reviewed_excerpt": excerpt,
                    "excerpt_kind": "authoring_summary",
                    "extraction_method": "feedback_workbook_transcription",
                    "review_status": "needs_review",
                    "fragment_checksum": "sha256:" + "0" * 64,
                }
            )
    raw = {
        "schema_version": "1.0.0",
        "documents": _document_rows(),
        "fragments": fragments,
        "manifest_checksum": "sha256:" + "0" * 64,
    }
    return finalize_source_manifest(raw)


def finalize_source_manifest(raw: dict[str, Any]) -> dict[str, Any]:
    result = deepcopy(raw)
    for document in result["documents"]:
        path = ROOT / document["file_path"]
        if not path.is_file():
            raise ValueError(f"来源 PDF 不存在: {document['file_path']}")
        document["pdf_checksum"] = sha256_digest(path.read_bytes())
    for fragment in result["fragments"]:
        fragment["fragment_checksum"] = fragment_payload_checksum(fragment)
    result["manifest_checksum"] = source_manifest_payload_checksum(result)
    SourceManifest.model_validate(result)
    return result


def _coverage_class(value: str) -> str:
    for grade in ("A", "B", "C"):
        if f"（{grade}）" in value:
            return grade
    raise ValueError(f"无法解析可行性等级: {value!r}")


def _evidence_domains(value: Any) -> list[str]:
    return [item.strip() for item in re.split(r"[/、]", str(value)) if item.strip()]


def _fact_type(node_type: str, evidence: str) -> str:
    if node_type == "SCORE":
        return "composite_score"
    if "检验日期" in evidence:
        return "observation_series"
    if "检验" in evidence:
        return "laboratory_observation"
    if "手术" in evidence:
        return "procedure_history"
    if any(term in evidence for term in ("检查", "心超", "ECG", "Holter", "X线", "肺功能", "电生理")):
        return "examination_finding"
    if "诊断" in evidence and "文书" not in evidence:
        return "diagnosis_history"
    return "clinical_assertion"


def _threshold(value: str) -> int | None:
    for pattern in (r"至少\s*(\d+)", r"≥\s*(\d+)"):
        match = re.search(pattern, value)
        if match:
            return int(match.group(1))
    return None


def _policy_hint(value: str, pattern: str) -> dict[str, Any]:
    if re.search(pattern, value):
        return {"source_declaration": value, "review_status": "needs_review"}
    return {"mode": "not_applicable_or_unspecified"}


def _criterion_node(row: dict[str, Any], child_count: int) -> dict[str, Any]:
    rule_id = str(row["规则ID"])
    local_id = str(row["节点ID"])
    source_type = str(row["节点类型"])
    condition = str(row["节点条件"])
    operator_text = str(row["运算符/阈值"] or "")
    evidence = str(row["证据域"])
    node_id = f"{rule_id}.{local_id}"
    parent = str(row["父节点ID"] or "").strip()
    blocked_root = rule_id in BLOCKERS and local_id == "ROOT"
    if blocked_root:
        node_type, operator, threshold, compilation = (
            "BLOCKED_ROOT",
            None,
            None,
            "blocked",
        )
    elif source_type in {"ALL", "ANY", "AT_LEAST_N"}:
        node_type = "LOGIC"
        operator = {"ALL": "AND", "ANY": "OR", "AT_LEAST_N": "AT_LEAST_N"}[
            source_type
        ]
        threshold = _threshold(operator_text) if operator == "AT_LEAST_N" else None
        if operator == "AT_LEAST_N" and threshold is None:
            raise ValueError(f"{node_id} 缺少 AT_LEAST_N threshold")
        # 反馈表只有“主要节点”；无显式 child 的逻辑摘要绝不冒充完整可执行树。
        compilation = "partial"
    else:
        node_type, operator, threshold = "LEAF", None, None
        compilation = "compiled" if str(row["机器化状态"]) == "READY" else "partial"
    refs = [
        _fragment_id(rule_id, local_id, page)
        for page in _page_numbers(str(row["2025 PDF页"]))
    ]
    return {
        "node_id": node_id,
        "parent_node_id": f"{rule_id}.{parent}" if parent else None,
        "node_type": node_type,
        "operator": operator,
        "threshold": threshold,
        "criterion_id": node_id if node_type == "LEAF" else "",
        "fact_type": _fact_type(source_type, evidence) if node_type == "LEAF" else "",
        "summary": condition,
        "expected_condition": {
            "summary": condition,
            "operator_or_threshold": operator_text,
            "source_node_type": source_type,
        },
        "evidence_domains": _evidence_domains(evidence),
        "unit_policy": _policy_hint(operator_text, r"g/L|mg/g|g/24h|ml/|mmHg|%|秒"),
        "repetition_policy": _policy_hint(operator_text + condition, r"\d+次|全时限|累计"),
        "temporal_policy": _policy_hint(operator_text + condition, r"\d+个月|\d+年|\d+天|时限"),
        "source_fragment_ids": refs,
        "compilation_status": compilation,
        "authoring_status": str(row["机器化状态"]),
        "notes": str(row["线下问题/备注"] or ""),
    }


def build_criteria_asset(
    disease_rows: list[dict[str, Any]],
    node_rows: list[dict[str, Any]],
    source_manifest: dict[str, Any],
) -> dict[str, Any]:
    grouped: dict[str, list[dict[str, Any]]] = {rule_id: [] for rule_id in DISEASE_IDS}
    for row in node_rows:
        grouped[str(row["规则ID"])].append(row)
    revisions = []
    for disease in disease_rows:
        rule_id = str(disease["规则ID（暂定）"])
        rows = grouped[rule_id]
        parent_counts: dict[str, int] = {}
        for row in rows:
            parent = str(row["父节点ID"] or "").strip()
            if parent:
                parent_counts[parent] = parent_counts.get(parent, 0) + 1
        nodes = [
            _criterion_node(row, parent_counts.get(str(row["节点ID"]), 0))
            for row in rows
        ]
        root = next(item for item in nodes if item["parent_node_id"] is None)
        blocker = None
        if rule_id in BLOCKERS:
            spec = BLOCKERS[rule_id]
            blocker = {
                "block_reason_code": spec["block_reason_code"],
                "unresolved_question": spec["unresolved_question"],
                "affected_node_ids": [f"{rule_id}.{item}" for item in spec["affected"]],
                "required_approver_roles": spec["roles"],
                "prohibited_fallbacks": ["2020_POLICY", "LLM", "LOCAL_DEFAULT"],
                "resolution_reference": None,
            }
        revisions.append(
            {
                "rule_id": rule_id,
                "canonical_disease_id": DISEASE_IDS[rule_id],
                "disease_name": str(disease["病种"]),
                "coverage_class": _coverage_class(
                    str(disease["是否可以覆盖（可行性判断）"])
                ),
                "revision_id": f"{rule_id}-2025-r1",
                "policy_version": "2025",
                "lifecycle": "draft",
                "review_status": "needs_review",
                "execution_status": "BLOCKED" if blocker else "EVALUATABLE",
                "effective_from": "2025-09-02",
                "effective_to": None,
                "root_node_id": root["node_id"],
                "source_fragment_ids": root["source_fragment_ids"],
                "nodes": nodes,
                "blocker": blocker,
                "coverage_notes": str(disease["备注"] or ""),
            }
        )
    raw = {
        "schema_version": "1.0.0",
        "asset_type": "chronic_disease_criteria",
        "release_id": "chronic-criteria-2025-draft-r1",
        "source_manifest_checksum": source_manifest["manifest_checksum"],
        "ordered_disease_revision_ids": [item["revision_id"] for item in revisions],
        "policy_sets": [
            {
                "policy_id": "hlj-outpatient-chronic-2025",
                "policy_version": "2025",
                "policy_role": "current_recognition",
                "release_id": "chronic-criteria-2025-draft-r1",
                "lifecycle": "draft",
                "review_status": "needs_review",
                "execution_enabled": False,
                "source_document_ids": [CURRENT_SOURCE_ID],
                "effective_from": "2025-09-02",
                "effective_to": None,
                "disease_revisions": revisions,
                "notes": "新认定唯一主口径；所有 revision 待逐页双人转录和专家审批",
            },
            {
                "policy_id": "heihe-outpatient-chronic-2020",
                "policy_version": "2020",
                "policy_role": "historical_reference",
                "release_id": "chronic-criteria-2020-reference-r1",
                "lifecycle": "draft",
                "review_status": "needs_review",
                "execution_enabled": False,
                "source_document_ids": [HISTORICAL_SOURCE_ID],
                "effective_from": "2020-05-15",
                "effective_to": None,
                "disease_revisions": [],
                "notes": "仅保留历史版本边界，未编译历史树；不得向2025提供节点、阈值或默认值",
            },
        ],
        "asset_checksum": "sha256:" + "0" * 64,
    }
    return finalize_criteria_asset(raw, source_manifest)


def finalize_criteria_asset(
    raw: dict[str, Any], source_manifest: dict[str, Any]
) -> dict[str, Any]:
    result = deepcopy(raw)
    result["source_manifest_checksum"] = source_manifest["manifest_checksum"]
    result["asset_checksum"] = asset_payload_checksum(result)
    asset = ChronicDiseaseCriteriaAsset.model_validate(result)
    manifest = SourceManifest.model_validate(source_manifest)
    validate_source_isolation(asset, manifest)
    return result


def schema_payload() -> dict[str, Any]:
    return ChronicDiseaseCriteriaAsset.model_json_schema()


def import_feedback(path: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    diseases, nodes = _feedback_rows(path)
    manifest = build_source_manifest(nodes)
    asset = build_criteria_asset(diseases, nodes, manifest)
    return manifest, asset


def _new_node(
    template: dict[str, Any], local_id: str, parent: str, summary: str,
    *, operator: str | None = None, condition: dict[str, Any] | None = None,
) -> dict[str, Any]:
    node = deepcopy(template)
    prefix = template["node_id"].split(".")[0]
    node.update(
        node_id=f"{prefix}.{local_id}", parent_node_id=f"{prefix}.{parent}",
        node_type="LOGIC" if operator else "LEAF", operator=operator, threshold=None,
        criterion_id="" if operator else f"{prefix}.{local_id}",
        fact_type="" if operator else "clinical_assertion",
        summary=summary, expected_condition=condition or {"summary": summary},
        compilation_status="partial", authoring_status="REVIEW",
        expert_interpretation_ids=[], notes="仅按已有2025摘要细化；原文和适用条件待复核",
    )
    if condition and condition.get("type") == "numeric":
        node["fact_type"] = "laboratory_observation"
        node["unit_policy"] = {"canonical_unit": condition["unit"], "conversion": "identity_only"}
    return node


def _apply_interpretation(revision: dict[str, Any], feedback: dict[str, Any]) -> None:
    rid = revision["rule_id"]
    revision["previous_revision_id"] = revision["revision_id"]
    revision["revision_id"] = f"{rid}-2025-r2"
    revision["expert_interpretation_ids"] = [feedback["interpretation_id"]]
    revision["review_status"] = "needs_review"
    revision["lifecycle"] = "draft"
    revision["coverage_notes"] += "；H列书面解释仅消除所列组合歧义，G列及原文审批仍待确认。"
    nodes = {n["node_id"].split(".")[1]: n for n in revision["nodes"]}
    root = nodes["ROOT"]
    # 其他病种仍只有主要节点摘要，不能因组合消歧变成完整树。
    for node in revision["nodes"]:
        if node["node_type"] != "BLOCKED_ROOT":
            node["compilation_status"] = "partial"
    if rid in {"CD01", "CD07", "CD09"}:
        revision["blocker"] = None
        revision["execution_status"] = "EVALUATABLE"
        root.update(node_type="LOGIC", operator="OR" if rid == "CD07" else "AND",
                    compilation_status="partial", authoring_status="REVIEW")
        root["summary"] = feedback["statement"] if rid != "CD01" else "临床表现 AND 诊断依据四项全部满足"
        root["expected_condition"] = {"summary": root["summary"]}
        root["notes"] = "组合关系由H列书面解释确认；其他节点与原文仍待复核"
    if rid == "CD01":
        nodes["DIAGNOSTIC_ALL"] = _new_node(root, "DIAGNOSTIC_ALL", "ROOT", feedback["statement"], operator="AND")
        for name in ("CBC_2OF3", "MARROW_ASP", "MARROW_BIOPSY", "EXCLUSION"):
            nodes[name]["parent_node_id"] = "CD01.DIAGNOSTIC_ALL"
            nodes[name]["notes"] = "H7明确诊断依据块内AND；具体阈值和原文仍待复核"
    elif rid == "CD07":
        nodes["GRADE12_DAMAGE"]["summary"] = "高血压1-2级 AND（靶器官损害 OR 临床综合征）"
        nodes["GRADE12_DAMAGE"]["expected_condition"] = {"summary": nodes["GRADE12_DAMAGE"]["summary"]}
        nodes["GRADE12_DAMAGE"]["notes"] = "H4明确组合；靶器官及综合征子项仍未完整转录"
    elif rid == "CD09":
        nodes["HISTORY"] = _new_node(root, "HISTORY", "ROOT", "慢性病毒性肝炎病史是必备条件")
        nodes["VIRAL_ANY"] = _new_node(root, "VIRAL_ANY", "ROOT", "HBV或HCV分支", operator="OR")
        for name in ("HBV_ANY", "HCV_ANY"):
            nodes[name]["parent_node_id"] = "CD09.VIRAL_ANY"
    elif rid == "CD19":
        nodes["SCORE_XRAY"] = _new_node(root, "SCORE_XRAY", "ROOT", feedback["statement"], operator="AND")
        for name in ("SCORE", "XRAY"):
            nodes[name]["parent_node_id"] = "CD19.SCORE_XRAY"
        nodes["SCORE"]["expected_condition"] = {
            "type": "numeric", "concept_id": "acr_eular_total_score", "terms": ["ACR/EULAR总分"],
            "operator": "gte", "value": 6, "unit": "分",
            "summary": "已记录的ACR/EULAR总分至少6分；评分分项未完整转录",
        }
        nodes["SCORE"]["notes"] = "阈值来自H5书面解释，未借用2020；不从未编译分项自动算分"
        nodes["SCORE"]["authoring_status"] = "REVIEW"
        nodes["XRAY"]["expected_condition"] = {"summary": "X线III期及以上", "operator_or_threshold": "III期及以上"}
        nodes["FUNCTION"] = _new_node(root, "FUNCTION", "ROOT", "严重肢体功能障碍分级（待专家明确）")
        root["notes"] = "H5已明确总分≥6且X线III期+；功能障碍分级仍未说明"
        revision["blocker"].update(
            block_reason_code="FUNCTIONAL_IMPAIRMENT_THRESHOLD_MISSING",
            unresolved_question="严重肢体功能障碍的分级及阈值是什么？H5未说明此项。",
            affected_node_ids=["CD19.ROOT", "CD19.FUNCTION"],
            resolution_reference=feedback["interpretation_id"],
        )
    elif rid == "CD03":
        dkd = nodes["DKD"]
        dkd["notes"] = "H8明确四个子项任一；数值沿用2025作者摘要，尚未全篇审批"
        dkd["expected_condition"]["operator_or_threshold"] = "四项任一（H8书面解释）"
        def numeric(name: str, parent: str, concept: str, op: str, value: float, unit: str) -> None:
            nodes[name] = _new_node(dkd, name, parent, f"{concept} {op} {value} {unit}", condition={
                "type": "numeric", "concept_id": concept, "operator": op, "value": value, "unit": unit, "terms": [concept],
            })
            nodes[name]["compilation_status"] = "compiled"
        numeric("DKD_EGFR", "DKD", "egfr", "lt", 60, "mL/min/1.73m2")
        nodes["DKD_PROTEIN_POSITIVE"] = _new_node(dkd, "DKD_PROTEIN_POSITIVE", "DKD", "尿蛋白阳性", condition={"type": "categorical", "concept_id": "urine_protein_qualitative", "terms": ["尿蛋白"], "values": ["+", "++", "+++", "++++", "阳性"]})
        nodes["DKD_PROTEIN_QUANT"] = _new_node(dkd, "DKD_PROTEIN_QUANT", "DKD", "尿蛋白定量两条单位路径任一", operator="OR")
        numeric("DKD_PROTEIN_RATIO", "DKD_PROTEIN_QUANT", "urine_protein_creatinine_ratio", "gte", 300, "mg/g")
        numeric("DKD_PROTEIN_24H", "DKD_PROTEIN_QUANT", "urine_protein_24h", "gte", 0.5, "g/24h")
        numeric("DKD_UACR", "DKD", "uacr", "gte", 300, "mg/g")
        for name, terms in {
            "DKD_EGFR": ["eGFR", "肾小球滤过率"],
            "DKD_PROTEIN_RATIO": ["尿蛋白/肌酐", "尿蛋白肌酐比"],
            "DKD_PROTEIN_24H": ["24小时尿蛋白", "24h尿蛋白"],
            "DKD_UACR": ["UACR", "尿白蛋白/肌酐", "尿白蛋白肌酐比"],
        }.items():
            nodes[name]["expected_condition"]["terms"] = terms
        nodes["DM_HISTORY"]["expected_condition"] = {"type": "presence", "concept_id": "diabetes_history", "terms": ["糖尿病"], "summary": "明确糖尿病病史"}
        for name in ("ROOT", "COMPLICATION_ANY", "DKD", "DKD_PROTEIN_QUANT", "DKD_PROTEIN_POSITIVE", "DM_HISTORY"):
            nodes[name]["compilation_status"] = "compiled"
        # NEUROPATHY/PAD/CAD_BRANCH仍为partial摘要，完整性闸不能被DKD局部展开绕过。
        nodes["COMPLICATION_ANY"]["notes"] = "H8仅明确肾病子项OR；其余并发症仍有partial节点"
    revision["nodes"] = list(nodes.values())
    for node_id in feedback["affected_node_ids"]:
        nodes[node_id.split(".")[1]]["expert_interpretation_ids"] = [feedback["interpretation_id"]]


def rebuild_existing() -> tuple[dict[str, Any], dict[str, Any]]:
    """从不可变r1和最小书面解释重建r2，不依赖原xlsx或当前生成文件。"""
    baseline_manifest = load_source_manifest(DRAFT_SOURCES, root=ROOT)
    load_criteria_asset(DRAFT_CRITERIA, source_manifest=baseline_manifest)
    manifest = json.loads(DRAFT_SOURCES.read_text(encoding="utf-8"))
    asset = json.loads(DRAFT_CRITERIA.read_text(encoding="utf-8"))
    feedback = json.loads(FEEDBACK_OUT.read_text(encoding="utf-8"))
    for item in feedback:
        ExpertInterpretation.model_validate(item)
        if item["interpretation_checksum"] != interpretation_payload_checksum(item):
            raise ValueError("expert interpretation checksum 不一致")
    if [(item["rule_id"], item["cell"], item["statement"]) for item in feedback] != R2_INTERPRETATIONS:
        raise ValueError("r2书面解释内容或坐标变化，必须创建新revision并调整编译逻辑")
    manifest["schema_version"] = "1.1.0"
    manifest["previous_manifest_checksum"] = manifest["manifest_checksum"]
    manifest["expert_interpretations"] = feedback
    manifest = finalize_source_manifest(manifest)
    policy = asset["policy_sets"][0]
    revisions = {item["rule_id"]: item for item in policy["disease_revisions"]}
    for item in feedback:
        _apply_interpretation(revisions[item["rule_id"]], item)
    asset["schema_version"] = "1.1.0"
    asset["previous_asset_checksum"] = asset["asset_checksum"]
    asset["release_id"] = policy["release_id"] = "chronic-criteria-2025-shadow-r2"
    asset["ordered_disease_revision_ids"] = [item["revision_id"] for item in policy["disease_revisions"]]
    policy["notes"] = "H4-H8书面解释已落盘；仅shadow，所有revision保持needs_review，原文与完整树审批未完成"
    return manifest, finalize_criteria_asset(asset, manifest)


def _write(manifest: dict[str, Any], asset: dict[str, Any]) -> None:
    SOURCES_OUT.write_bytes(pretty_json_bytes(manifest))
    CRITERIA_OUT.write_bytes(pretty_json_bytes(asset))
    SCHEMA_OUT.write_bytes(pretty_json_bytes(schema_payload()))
    SOURCE_SCHEMA_OUT.write_bytes(pretty_json_bytes(SourceManifest.model_json_schema()))


def _check(manifest: dict[str, Any], asset: dict[str, Any]) -> None:
    expected = {
        SOURCES_OUT: pretty_json_bytes(manifest),
        CRITERIA_OUT: pretty_json_bytes(asset),
        SCHEMA_OUT: pretty_json_bytes(schema_payload()),
        SOURCE_SCHEMA_OUT: pretty_json_bytes(SourceManifest.model_json_schema()),
    }
    stale = [str(path.relative_to(ROOT)) for path, data in expected.items() if path.read_bytes() != data]
    if stale:
        raise SystemExit(f"慢病知识资产不是确定性最新构建: {', '.join(stale)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--import-feedback", type=Path)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    if args.import_feedback:
        manifest, asset = import_feedback(args.import_feedback)
    else:
        manifest, asset = rebuild_existing()
    if args.check:
        _check(manifest, asset)
    else:
        _write(manifest, asset)
    print(
        json.dumps(
            {
                "diseases": 20,
                "nodes": sum(
                    len(item["nodes"])
                    for item in asset["policy_sets"][0]["disease_revisions"]
                ),
                "source_fragments": len(manifest["fragments"]),
                "asset_checksum": asset["asset_checksum"],
                "status": "checked" if args.check else "written",
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
